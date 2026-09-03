"""Regression tests for the cleaning pipeline.

The production hang happened AFTER «Archivo estructurado con éxito», while
generate_output_excel walked every cell with ws.max_row (O(n²)) and kept a
second full openpyxl workbook in RAM.

Run:
    python3 -m unittest tests.test_pipeline -v

To exercise a ~10–17MB xlsx locally (not part of the default unittest run):
    python3 scripts/verify_large_xlsx.py --target-mb 12
"""
from __future__ import annotations

import io
import time
import unittest

from openpyxl import Workbook, load_workbook

from pipeline import (
    KEY_MAP,
    OUTPUT_COLUMNS,
    detectar_duplicados_avanzado,
    extract_hyperlinks_from_xlsx,
    generate_output_excel,
    process_dossier,
)


HEADERS = [
    "NoticiaId", "Fecha", "Hora", "Medio", "Tipo de Medio", "Sección - Programa",
    "Título", "Autor - Conductor", "Nro. Pagina", "Dimensioncm2", "Duración - Nro. Caracteres",
    "CPE", "Valor de Nota", "Tier", "Audiencia", "Tono", "Tematica", "Subtema", "Producto",
    "Tipo de información", "Nombre vocero", "Mención en Titulo", "Mención en Foto",
    "Tipo mencion", "Tipo mencion 2", "Aparece Logo", "Resumen - Aclaracion", "CuerpoEs",
    "URL Nota AV", "URL (Streaming - Imagen)", "URL Nota", "Menciones - Empresa", "Empresa rel.",
]


def _body(i: int, length: int = 400) -> str:
    extra = " ".join(f"w{i}_{j}" for j in range(max(8, length // 8)))
    return (f"Artículo {i}. Párrafo <br> con texto. " + extra)[:length]


def build_dossier_xlsx(
    n_rows: int,
    body_len: int = 400,
    with_hyperlinks: bool = True,
    duplicate_url_every: int = 0,
    extra_mencion: bool = False,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(HEADERS)
    url_col = HEADERS.index("URL Nota") + 1
    stream_col = HEADERS.index("URL (Streaming - Imagen)") + 1

    for i in range(n_rows):
        tipo = ["Internet", "Radio", "Televisión", "Prensa"][i % 4]
        medio = ["El Tiempo", "Caracol", "RCN", "Semana"][i % 4]
        url = f"https://example.com/nota/{i}"
        if duplicate_url_every and i >= duplicate_url_every and i % duplicate_url_every == 0:
            url = f"https://example.com/nota/{i - duplicate_url_every}"
        body = _body(i, body_len)
        menciones = "Empresa A; Empresa B" if extra_mencion and i == 0 else "Empresa A"
        row = {
            "NoticiaId": 100000 + i,
            "Fecha": "15/01/2024",
            "Hora": "10:00:00",
            "Medio": medio,
            "Tipo de Medio": tipo,
            "Sección - Programa": "Política",
            "Título": f"Gobierno anuncia medida {i} | {medio}",
            "Autor - Conductor": "Redacción",
            "Nro. Pagina": 12,
            "Dimensioncm2": 150.5,
            "Duración - Nro. Caracteres": 45,
            "CPE": 2500000,
            "Valor de Nota": 1800000,
            "Tier": 1,
            "Audiencia": 120000,
            "Tono": "Neutro",
            "Tematica": "Política",
            "Subtema": "Gobierno",
            "Producto": "Marca",
            "Tipo de información": "Noticia",
            "Nombre vocero": "vocero",
            "Mención en Titulo": "Si",
            "Mención en Foto": "No",
            "Tipo mencion": "Directa",
            "Tipo mencion 2": "",
            "Aparece Logo": "Si",
            "Resumen - Aclaracion": body[:180],
            "CuerpoEs": body,
            "URL Nota AV": f"https://example.com.ar/av/{i}",
            "URL (Streaming - Imagen)": f"https://example.com/img/{i}",
            "URL Nota": "Link" if with_hyperlinks else url,
            "Menciones - Empresa": menciones,
            "Empresa rel.": "Empresa A; Empresa B" if extra_mencion and i == 0 else "Empresa A",
        }
        ws.append([row[h] for h in HEADERS])
        excel_row = i + 2
        if with_hyperlinks:
            cell = ws.cell(row=excel_row, column=url_col)
            cell.hyperlink = url
            cell.value = "Link"
            sc = ws.cell(row=excel_row, column=stream_col)
            sc.hyperlink = row["URL (Streaming - Imagen)"]
            sc.value = "Link"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


REGION_MAP = {
    "el tiempo": "Bogotá",
    "caracol": "Nacional",
    "rcn": "Nacional",
    "semana": "Bogotá",
}
INTERNET_MAP = {"el tiempo": "El Tiempo", "semana": "Semana"}


class TestHyperlinks(unittest.TestCase):
    def test_extract_hyperlinks_from_xlsx(self):
        data = build_dossier_xlsx(5, body_len=80, with_hyperlinks=True)
        links = extract_hyperlinks_from_xlsx(data, "Datos")
        url_col = HEADERS.index("URL Nota") + 1
        self.assertIn((2, url_col), links)
        self.assertEqual(links[(2, url_col)], "https://example.com/nota/0")


class TestDuplicates(unittest.TestCase):
    def _row(self, **kwargs):
        base = {
            "ID Noticia": 1,
            "Tipo de Medio": "Internet",
            "Medio": "El Tiempo",
            "Título": "Titular",
            "Hora": "10:00:00",
            "URL Nota": None,
            "Link Nota": None,
            "Link (Streaming - Imagen)": None,
            "Menciones - Empresa": "Acme",
            "is_duplicate": False,
        }
        base.update(kwargs)
        return base

    def test_same_url_nota_and_mencion_is_duplicate_regardless_of_title(self):
        rows = [
            self._row(
                **{
                    "ID Noticia": 1,
                    "Título": "Alpha titular único",
                    "URL Nota": {"value": "Link", "url": "https://www.example.com/a"},
                    "Link (Streaming - Imagen)": {"value": "Link", "url": "https://www.example.com/a"},
                }
            ),
            self._row(
                **{
                    "ID Noticia": 2,
                    "Título": "Beta titular distinto por completo xyz",
                    "URL Nota": {"value": "Link", "url": "https://example.com/a"},
                    "Link (Streaming - Imagen)": {"value": "Link", "url": "https://example.com/a"},
                }
            ),
        ]
        out = detectar_duplicados_avanzado(rows, KEY_MAP)
        self.assertFalse(out[0]["is_duplicate"])
        self.assertTrue(out[1]["is_duplicate"])
        self.assertEqual(out[1]["ID duplicada"], 1)

    def test_same_mencion_different_url_nota_is_not_duplicate(self):
        rows = [
            self._row(
                **{
                    "ID Noticia": 10,
                    "Título": "Gobierno anuncia reforma tributaria nacional",
                    "URL Nota": {"value": "Link", "url": "https://example.com/t1"},
                    "Link (Streaming - Imagen)": {"value": "Link", "url": "https://example.com/t1"},
                }
            ),
            self._row(
                **{
                    "ID Noticia": 11,
                    "Título": "Gobierno anuncia reforma tributaria nacional hoy",
                    "URL Nota": {"value": "Link", "url": "https://example.com/t2"},
                    "Link (Streaming - Imagen)": {"value": "Link", "url": "https://example.com/t2"},
                }
            ),
        ]
        out = detectar_duplicados_avanzado(rows, KEY_MAP)
        self.assertFalse(out[0]["is_duplicate"])
        self.assertFalse(out[1]["is_duplicate"])

    def test_av_same_mencion_medio_hora_is_duplicate(self):
        for tipo in ("Radio", "Televisión", "AM", "FM", "Aire", "Cable"):
            with self.subTest(tipo=tipo):
                rows = [
                    self._row(
                        **{
                            "ID Noticia": 1,
                            "Tipo de Medio": tipo,
                            "Medio": "Caracol",
                            "Hora": "10:00",
                            "Título": "Noticiero mediodía",
                            "Menciones - Empresa": "Acme",
                        }
                    ),
                    self._row(
                        **{
                            "ID Noticia": 2,
                            "Tipo de Medio": tipo,
                            "Medio": "Caracol",
                            "Hora": "10:00:00",
                            "Título": "Otro título totalmente distinto",
                            "Menciones - Empresa": "Acme",
                        }
                    ),
                ]
                out = detectar_duplicados_avanzado(rows, KEY_MAP)
                self.assertFalse(out[0]["is_duplicate"])
                self.assertTrue(out[1]["is_duplicate"], tipo)
                self.assertEqual(out[1]["ID duplicada"], 1)

    def test_av_different_hora_is_not_duplicate_even_with_same_title(self):
        rows = [
            self._row(
                **{
                    "ID Noticia": 1,
                    "Tipo de Medio": "Radio",
                    "Medio": "Caracol",
                    "Hora": "10:00:00",
                    "Título": "Gobierno anuncia reforma tributaria nacional",
                    "Menciones - Empresa": "Acme",
                }
            ),
            self._row(
                **{
                    "ID Noticia": 2,
                    "Tipo de Medio": "Radio",
                    "Medio": "Caracol",
                    "Hora": "11:00:00",
                    "Título": "Gobierno anuncia reforma tributaria nacional",
                    "Menciones - Empresa": "Acme",
                }
            ),
        ]
        out = detectar_duplicados_avanzado(rows, KEY_MAP)
        self.assertFalse(out[0]["is_duplicate"])
        self.assertFalse(out[1]["is_duplicate"])


class TestExport(unittest.TestCase):
    def test_write_only_export_preserves_links_and_dates(self):
        import datetime
        import pandas as pd

        rows = [{
            "ID Noticia": 1,
            "Fecha": pd.Timestamp("2024-01-15"),
            "Hora": "10:00",
            "Medio": "El Tiempo",
            "Tipo de Medio": "Internet",
            "Sección - Programa": "Política",
            "Región": "Bogotá",
            "Título": "Titulo de prueba | Medio",
            "Autor - Conductor": "Autor",
            "Nro. Pagina": 12,
            "Dimensión": 100,
            "Duración - Nro. Caracteres": 2000,
            "CPE": 1500000,
            "Tier": 1,
            "Audiencia": 50000,
            "Tono": "Neutro",
            "Tema": "Tema",
            "Subtema": "Sub",
            "Producto": "Prod",
            "Tipo de información": "Info",
            "Nombre vocero": "Vocero",
            "Mención en Titulo": "Si",
            "Mención en Foto": "No",
            "Tipo mencion": "Directa",
            "Tipo mencion 2": "",
            "Aparece Logo": "Si",
            "revalorización": 1000,
            "resumen corto": "Resumen",
            "Link Nota": {"value": "Link", "url": "https://example.com/nota/1"},
            "Resumen - Aclaracion": "Cuerpo de la nota.",
            "Link (Streaming - Imagen)": {"value": "Link", "url": "https://example.com/img/1"},
            "Menciones - Empresa": "Empresa",
            "ID duplicada": "",
            "is_duplicate": False,
        }]
        data = generate_output_excel(rows, KEY_MAP)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        self.assertEqual(ws.title, "Resultado")
        self.assertEqual([c.value for c in ws[1]], OUTPUT_COLUMNS)
        self.assertEqual(ws.cell(2, 1).value, 1)
        fecha = ws.cell(2, OUTPUT_COLUMNS.index("Fecha") + 1)
        self.assertIsInstance(fecha.value, datetime.datetime)
        self.assertEqual(fecha.number_format, "DD/MM/YYYY")
        link = ws.cell(2, OUTPUT_COLUMNS.index("Link Nota") + 1)
        self.assertEqual(link.value, "Link")
        self.assertEqual(link.hyperlink.target, "https://example.com/nota/1")

    def test_export_many_rows_finishes_quickly(self):
        """Catches the old O(n²) ws.max_row-per-row export hang."""
        import inspect
        from pipeline import generate_output_excel as _export
        self.assertNotIn(
            "ws.max_row",
            inspect.getsource(_export),
            "export must not call ws.max_row per row (quadratic on large sheets)",
        )
        n = 4000
        rows = []
        for i in range(n):
            rows.append({
                "ID Noticia": i,
                "Fecha": None,
                "Hora": "10:00",
                "Medio": "El Tiempo",
                "Tipo de Medio": "Internet",
                "Título": f"Titulo {i}",
                "CPE": 1000,
                "Link Nota": {"value": "Link", "url": f"https://example.com/{i}"},
                "Resumen - Aclaracion": "x" * 80,
                "Menciones - Empresa": "Empresa",
                "is_duplicate": False,
            })
        t0 = time.time()
        data = generate_output_excel(rows, KEY_MAP)
        elapsed = time.time() - t0
        self.assertGreater(len(data), 1000)
        self.assertLess(
            elapsed,
            20.0,
            f"export of {n} rows took {elapsed:.1f}s; the old max_row loop would grow quadratically",
        )
        wb = load_workbook(io.BytesIO(data))
        try:
            self.assertEqual(wb.active.max_row, n + 1)
        finally:
            wb.close()


class TestProcessDossier(unittest.TestCase):
    def test_pipeline_completes_and_reports_export_progress(self):
        data = build_dossier_xlsx(
            80,
            body_len=220,
            with_hyperlinks=True,
            duplicate_url_every=20,
            extra_mencion=True,
        )
        stages = []

        def progress(pct, msg):
            stages.append((pct, msg))

        result = process_dossier(io.BytesIO(data), REGION_MAP, INTERNET_MAP, progress=progress)
        self.assertTrue(result["total_rows"] >= 80)
        # extra_mencion on row 0 duplicates Empresa A; Empresa B -> +1 row
        self.assertGreaterEqual(result["total_rows"], 81)
        self.assertGreater(result["unique_rows"], 0)
        self.assertGreater(len(result["output_data"]), 1000)
        messages = " | ".join(m for _, m in stages)
        self.assertIn("Archivo estructurado con éxito", messages)
        self.assertIn("Generando archivo de resultado", messages)
        self.assertIn("Limpieza completada", messages)
        self.assertEqual(stages[-1][0], 100)

        wb = load_workbook(io.BytesIO(result["output_data"]))
        ws = wb.active
        self.assertEqual(ws.max_row, result["total_rows"] + 1)
        self.assertEqual(ws.max_column, len(OUTPUT_COLUMNS))

    def test_unmapped_medio_is_reported(self):
        data = build_dossier_xlsx(4, body_len=40, with_hyperlinks=False)
        result = process_dossier(io.BytesIO(data), {}, {}, progress=None)
        self.assertTrue(result["medios_sin_mapear"])


if __name__ == "__main__":
    unittest.main()
