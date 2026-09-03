#!/usr/bin/env python3
"""Build a large-ish unique-text xlsx and run the cleaning pipeline.

This is the practical check for the ~17MB production hang: after
«Archivo estructurado con éxito» the export used to stall with no further
UI. This script prints every pipeline stage (including export) and fails if
the run exceeds --timeout seconds.

Example (similar size to the reported file):
    python3 scripts/verify_large_xlsx.py --target-mb 17 --timeout 180
"""
from __future__ import annotations

import argparse
import io
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook  # noqa: E402

from pipeline import process_dossier  # noqa: E402
from tests.test_pipeline import HEADERS, REGION_MAP, INTERNET_MAP  # noqa: E402


def build_unique_xlsx(n_rows: int, body_len: int) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Datos")
    ws.append(HEADERS)
    for i in range(n_rows):
        tipo = ["Internet", "Radio", "Televisión", "Prensa"][i % 4]
        medio = ["El Tiempo", "Caracol", "RCN", "Semana"][i % 4]
        extra = " ".join(f"w{i}_{j}" for j in range(max(8, body_len // 8)))
        body = (f"Artículo {i}. " + extra)[:body_len]
        url = f"https://example.com/nota/{i}"
        row = {
            "NoticiaId": 100000 + i,
            "Fecha": "15/01/2024",
            "Hora": "10:00:00",
            "Medio": medio,
            "Tipo de Medio": tipo,
            "Sección - Programa": "Política",
            "Título": f"Gobierno anuncia medida {i} | {medio}",
            "Autor - Conductor": "Redacción",
            "Nro. Pagina": 12,
            "Dimensioncm2": 150.5,
            "Duración - Nro. Caracteres": 45,
            "CPE": 2500000,
            "Valor de Nota": 1800000,
            "Tier": 1,
            "Audiencia": 120000,
            "Tono": "Neutro",
            "Tematica": "Política",
            "Subtema": "Gobierno",
            "Producto": "Marca",
            "Tipo de información": "Noticia",
            "Nombre vocero": "vocero",
            "Mención en Titulo": "Si",
            "Mención en Foto": "No",
            "Tipo mencion": "Directa",
            "Tipo mencion 2": "",
            "Aparece Logo": "Si",
            "Resumen - Aclaracion": body[:180],
            "CuerpoEs": body,
            "URL Nota AV": f"https://example.com.ar/av/{i}",
            "URL (Streaming - Imagen)": f"https://example.com/img/{i}",
            "URL Nota": url,
            "Menciones - Empresa": "Empresa A",
            "Empresa rel.": "Empresa A",
        }
        ws.append([row[h] for h in HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target-mb", type=float, default=8.0)
    parser.add_argument("--timeout", type=float, default=120.0)
    parser.add_argument("--rows", type=int, default=0, help="Override row count (0 = derive from size)")
    parser.add_argument("--body-len", type=int, default=2200)
    args = parser.parse_args()

    n_rows = args.rows
    if n_rows <= 0:
        # Unique tokens compress poorly; ~2.2KB body * N rows ≈ target after zip.
        n_rows = max(1500, int(args.target_mb * 900))

    print(f"Building synthetic dossier: rows={n_rows} body_len={args.body_len}…")
    t_build = time.time()
    data = build_unique_xlsx(n_rows, args.body_len)
    size_mb = len(data) / (1024 * 1024)
    print(f"Built {size_mb:.2f} MB in {time.time() - t_build:.1f}s")

    if size_mb < args.target_mb * 0.6:
        print(
            f"File is smaller than requested ({size_mb:.2f} MB < {args.target_mb} MB); "
            "bump --rows / --body-len if you need a closer match."
        )

    stages = []

    def progress(pct, msg):
        line = f"[{pct:3d}%] {msg}"
        stages.append(line)
        print(line, flush=True)

    t0 = time.time()
    result = process_dossier(io.BytesIO(data), REGION_MAP, INTERNET_MAP, progress=progress)
    elapsed = time.time() - t0
    print(
        f"Done in {elapsed:.2f}s → {result['total_rows']} rows "
        f"({result['unique_rows']} unique, {result['duplicates']} dups), "
        f"output {len(result['output_data']) / 1024:.1f} KB"
    )

    joined = "\n".join(stages)
    if "Archivo estructurado con éxito" not in joined:
        print("FAIL: missing structured-file stage", file=sys.stderr)
        return 1
    if "Generando archivo de resultado" not in joined:
        print("FAIL: export never reported progress after structuring", file=sys.stderr)
        return 1
    if elapsed > args.timeout:
        print(f"FAIL: exceeded timeout {args.timeout}s", file=sys.stderr)
        return 1

    wb = load_workbook(io.BytesIO(result["output_data"]), read_only=True)
    try:
        out_rows = sum(1 for _ in wb.active.iter_rows(min_row=2))
    finally:
        wb.close()
    if out_rows != result["total_rows"]:
        print(f"FAIL: output rows {out_rows} != {result['total_rows']}", file=sys.stderr)
        return 1
    print("OK: pipeline continued after structuring and wrote the result xlsx.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
